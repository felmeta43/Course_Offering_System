"""Excel import/export helpers for bulk Course and StudentGroup data entry.

Admins fill in a spreadsheet (downloadable via the template generators
below) and upload it; we parse it here and create/update records. This is
what makes onboarding a whole college's curriculum and class lists fast.
"""
import openpyxl
from django.http import HttpResponse

from institutions.models import Department

COURSE_HEADERS = [
    "department_code",
    "course_code",
    "course_title",
    "credit_hours",
    "lecture_hours",
    "tutorial_hours",
    "lab_hours",
    "course_type",
    "year_level",
    "semester",
    "target_department_codes",
    "prerequisite_codes",
]

COURSE_SAMPLE_ROW = [
    "CS", "CS3101", "Emerging Technology", 3, 2, 0, 3, "COMMON", 3, 1, "EE,ME,CE", "",
]

STUDENT_GROUP_HEADERS = [
    "department_code",
    "academic_year",
    "semester",
    "year_level",
    "section",
    "number_of_students",
    "is_extension",
]

STUDENT_GROUP_SAMPLE_ROW = [
    "CS", "2025/2026", 1, 2, "A", 45, "NO",
]


def _build_template(headers, sample_row, sheet_title):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)
    ws.append(sample_row)
    return wb


def course_import_template_response():
    wb = _build_template(COURSE_HEADERS, COURSE_SAMPLE_ROW, "Courses")
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = "attachment; filename=course_import_template.xlsx"
    wb.save(response)
    return response


def student_group_import_template_response():
    wb = _build_template(STUDENT_GROUP_HEADERS, STUDENT_GROUP_SAMPLE_ROW, "StudentGroups")
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = "attachment; filename=student_group_import_template.xlsx"
    wb.save(response)
    return response


def _dept(code):
    code = (code or "").strip()
    if not code:
        return None
    return Department.objects.filter(code__iexact=code).first()


def import_courses_from_workbook(workbook_file):
    """Returns (created_count, updated_count, errors: list[str])."""
    from .models import Course

    wb = openpyxl.load_workbook(workbook_file, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    created, updated, errors = 0, 0, []

    # First pass: create/update courses without prerequisites/targets so all
    # course codes exist before we wire up the relational fields.
    pending_links = []
    for idx, row in enumerate(rows, start=2):
        if row is None or all(cell in (None, "") for cell in row):
            continue
        try:
            (dept_code, code, title, credit_hours, lecture_hours, tutorial_hours,
             lab_hours, course_type, year_level, semester, target_codes, prereq_codes) = row[:12]
        except ValueError:
            errors.append(f"Row {idx}: expected {len(COURSE_HEADERS)} columns")
            continue

        department = _dept(dept_code)
        if department is None:
            errors.append(f"Row {idx}: unknown department_code '{dept_code}'")
            continue
        if not code:
            errors.append(f"Row {idx}: missing course_code")
            continue

        course_type_norm = (str(course_type).strip().upper() if course_type else "MAJOR")
        if course_type_norm not in Course.CourseType.values:
            errors.append(f"Row {idx}: invalid course_type '{course_type}'")
            continue

        defaults = dict(
            department=department,
            title=title or code,
            credit_hours=credit_hours or 0,
            lecture_hours=lecture_hours or 0,
            tutorial_hours=tutorial_hours or 0,
            lab_hours=lab_hours or 0,
            course_type=course_type_norm,
            year_level=int(year_level) if year_level else 1,
            semester=int(semester) if semester else 1,
        )
        course, was_created = Course.objects.update_or_create(code=str(code).strip(), defaults=defaults)
        created += int(was_created)
        updated += int(not was_created)
        pending_links.append((course, target_codes, prereq_codes, idx))

    for course, target_codes, prereq_codes, idx in pending_links:
        if target_codes:
            target_depts = []
            for c in str(target_codes).split(","):
                d = _dept(c)
                if d:
                    target_depts.append(d)
                elif c.strip():
                    errors.append(f"Row {idx}: unknown target department '{c.strip()}'")
            course.target_departments.set(target_depts)
        if prereq_codes:
            prereq_courses = []
            for c in str(prereq_codes).split(","):
                c = c.strip()
                if not c:
                    continue
                pc = Course.objects.filter(code__iexact=c).first()
                if pc:
                    prereq_courses.append(pc)
                else:
                    errors.append(f"Row {idx}: unknown prerequisite course '{c}'")
            course.prerequisites.set(prereq_courses)

    return created, updated, errors


def import_student_groups_from_workbook(workbook_file, institution):
    """Returns (created_count, updated_count, errors: list[str])."""
    from institutions.models import AcademicTerm

    from .models import StudentGroup

    wb = openpyxl.load_workbook(workbook_file, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    created, updated, errors = 0, 0, []

    for idx, row in enumerate(rows, start=2):
        if row is None or all(cell in (None, "") for cell in row):
            continue
        try:
            (dept_code, academic_year, semester, year_level, section,
             number_of_students, is_extension) = row[:7]
        except ValueError:
            errors.append(f"Row {idx}: expected {len(STUDENT_GROUP_HEADERS)} columns")
            continue

        department = _dept(dept_code)
        if department is None:
            errors.append(f"Row {idx}: unknown department_code '{dept_code}'")
            continue

        term, _ = AcademicTerm.objects.get_or_create(
            institution=institution,
            academic_year=str(academic_year).strip(),
            semester=int(semester) if semester else 1,
        )

        is_ext_norm = str(is_extension).strip().upper() in ("YES", "TRUE", "1", "Y")

        defaults = dict(number_of_students=int(number_of_students) if number_of_students else 0)
        group, was_created = StudentGroup.objects.update_or_create(
            department=department,
            academic_term=term,
            year_level=int(year_level) if year_level else 1,
            section=str(section).strip() if section else "A",
            is_extension=is_ext_norm,
            defaults=defaults,
        )
        created += int(was_created)
        updated += int(not was_created)

    return created, updated, errors
